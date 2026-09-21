"""Independently authored reader tests; no organizer answers or generators."""
import hashlib
import importlib.util
import io
import unittest
from unittest.mock import patch

from harborlight.readers import FIELDS, _ocr_image, read_document


SI = """SHIPPING INSTRUCTION
Shipper/Exporter: Atlas Paper Ltd
  11 Harbor Road
  Singapore 048624
Consignee (Non-Negotiable): Beacon Trading
  15 Wharf Street
  London, United Kingdom
Notify Party/Intermediate Consignee: Beacon Logistics
Port of Loading (POL): SINGAPORE (SGSIN)
Discharge Port: FELIXSTOWE, UK (GBFXT)
No. of Containers or Packages: 2 x 40'HC
Gross Weight毛重(KGS): 48,200 KG
Vessel: MERIDIAN
Booking Reference: TEST-001
"""


class ReaderTests(unittest.TestCase):
    def test_text_fields_have_source_evidence(self):
        data = SI.encode()
        document = read_document(data, "sample.txt")
        self.assertEqual(document["kind"], "SI")
        self.assertIsNone(document["error"])
        self.assertEqual(set(document["fields"]), set(FIELDS))
        self.assertEqual(document["sha256"], hashlib.sha256(data).hexdigest())
        self.assertEqual(document["fields"]["gross_weight_kg"]["value"], "48,200 KG")
        self.assertEqual(document["fields"]["notify_party"]["value"], "Beacon Logistics")
        self.assertIn("11 Harbor Road", document["fields"]["shipper"]["value"])
        self.assertNotIn("Vessel", document["fields"]["gross_weight_kg"]["value"])
        for item in document["fields"].values():
            self.assertTrue(item["evidence"])
            self.assertEqual(item["page"], 1)
            self.assertGreater(item["confidence"], .9)

    def test_invoice_masquerading_as_bl_is_other(self):
        document = read_document(("COMMERCIAL INVOICE\nBill of Lading No.: TEST\n" + SI).encode(), "official_BL.txt")
        self.assertEqual(document["kind"], "OTHER")

    def test_filename_is_not_document_identity(self):
        document = read_document(b"Shipper: Atlas\nConsignee: Beacon", "BL.txt")
        self.assertEqual(document["kind"], "UNKNOWN")

    def test_bill_of_lading_instruction_is_si(self):
        document = read_document(SI.replace("SHIPPING INSTRUCTION", "BILL OF LADING INSTRUCTION\nB/L NUMBER: X").encode(), "BL.txt")
        self.assertEqual(document["kind"], "SI")

    def test_order_of_label_stops_shipper_address(self):
        text = SI.replace("Consignee (Non-Negotiable)", "To the Order of")
        document = read_document(text.encode(), "order.txt")
        self.assertTrue(document["fields"]["consignee"]["value"].startswith("Beacon Trading"))
        self.assertNotIn("Beacon", document["fields"]["shipper"]["value"])

    def test_placeholders_are_missing_and_not_fake_values(self):
        for placeholder in ("N/A", "—", "NOT PROVIDED", "TBD", "[blank]", ""):
            document = read_document(f"BILL OF LADING\nNotify Party: {placeholder}\nGross Weight: 12 KG".encode(), "a.txt")
            self.assertIsNone(document["fields"]["notify_party"]["value"], placeholder)
            self.assertEqual(document["fields"]["notify_party"]["confidence"], 0)

    def test_conflicting_repeated_fields_never_choose_silently(self):
        document = read_document((SI + "\nGross Weight: 50,000 KG").encode(), "a.txt")
        field = document["fields"]["gross_weight_kg"]
        self.assertIsNone(field["value"])
        self.assertTrue(field["ambiguous"])
        self.assertEqual(len(field["candidates"]), 2)
        self.assertTrue(any("Conflicting" in x for x in document["warnings"]))

    def test_repeated_identical_fields_are_not_conflicts(self):
        document = read_document((SI + "\nGross Weight: 48,200 KG").encode(), "a.txt")
        self.assertEqual(document["fields"]["gross_weight_kg"]["value"], "48,200 KG")
        self.assertEqual(document["fields"]["gross_weight_kg"]["occurrences"], 2)

    def test_explicit_total_wins_over_item_weight(self):
        text = "BILL OF LADING\nGross Weight: 12,000 KG\nTOTAL Gross Weight: 24,000 KG\n"
        document = read_document(text.encode(), "a.txt")
        self.assertEqual(document["fields"]["gross_weight_kg"]["value"], "24,000 KG")

    def test_weight_unit_in_label_is_preserved(self):
        document = read_document(b"BILL OF LADING\nGross Weight (MT): 48.2", "a.txt")
        self.assertEqual(document["fields"]["gross_weight_kg"]["value"], "48.2 MT")
        self.assertIn("(MT): 48.2", document["fields"]["gross_weight_kg"]["evidence"])

    def test_table_weight_header_does_not_capture_container_id(self):
        text = "BILL OF LADING\nCONTAINER NO.       DESCRIPTION       GROSS WEIGHT (KG)\nABCD1234567         40HC PAPER         12,000\nTotal Gross Weight: 12,000 KG\n"
        document = read_document(text.encode(), "a.txt")
        self.assertEqual(document["fields"]["gross_weight_kg"]["value"], "12,000 KG")

    def test_two_column_party_addresses(self):
        text = "SHIPPING INSTRUCTION\n" + "Shipper:".ljust(40) + "Consignee:\n" + "Atlas Ltd".ljust(40) + "Beacon Ltd\n" + "11 Harbor Road".ljust(40) + "15 Wharf Street\n" + "POL: SINGAPORE\nPOD: FELIXSTOWE\n"
        document = read_document(text.encode(), "a.txt")
        self.assertEqual(document["fields"]["shipper"]["value"], "Atlas Ltd\n11 Harbor Road")
        self.assertEqual(document["fields"]["consignee"]["value"], "Beacon Ltd\n15 Wharf Street")

    def test_utf16_text(self):
        document = read_document(SI.encode("utf-16"), "a.txt")
        self.assertIsNone(document["error"])
        self.assertEqual(document["kind"], "SI")

    def test_corrupt_pdf_is_visible_error(self):
        document = read_document(b"%PDF-1.7\nbroken", "BL.pdf")
        self.assertTrue(document["error"])
        self.assertEqual(document["kind"], "UNKNOWN")

    def test_empty_and_unsupported_content_fail_visibly(self):
        self.assertTrue(read_document(b"", "empty.txt")["error"])
        self.assertTrue(read_document(b"MZbinary", "BL.exe")["error"])

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx unavailable")
    def test_docx_bilingual_labels_tables_and_addresses(self):
        from docx import Document
        document = Document()
        document.add_paragraph("BILL OF LADING (DRAFT)")
        table = document.add_table(rows=0, cols=2)
        for label, value in [("Shipper (Principal or Seller) (发货人)", "Atlas Ltd\n11 Harbor Road"), ("Consignee (收货人)", "Beacon Ltd"), ("Notify (通知人)", "Beacon Logistics"), ("POL (装货港)", "SINGAPORE"), ("POD (卸货港)", "FELIXSTOWE"), ("Total Containers (箱数)", "2 x 40HC"), ("Gross Wt (kgs) (毛重 KGS)", "48,200")]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
        stream = io.BytesIO()
        document.save(stream)
        result = read_document(stream.getvalue(), "document.docx")
        self.assertIsNone(result["error"])
        self.assertEqual(result["kind"], "BL")
        self.assertTrue(all(item["value"] for item in result["fields"].values()))
        self.assertEqual(result["fields"]["shipper"]["value"], "Atlas Ltd\n11 Harbor Road")
        self.assertIn("logical", result["fields"]["shipper"]["locator"])

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx unavailable")
    def test_docx_parallel_party_columns_remain_separate(self):
        from docx import Document
        document = Document()
        document.add_paragraph("BILL OF LADING")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Shipper", "Consignee"
        table.cell(1, 0).text = "Atlas Ltd\n11 Harbor Road"
        table.cell(1, 1).text = "Beacon Ltd\n15 Wharf Street"
        stream = io.BytesIO()
        document.save(stream)
        result = read_document(stream.getvalue(), "columns.docx")
        self.assertEqual(result["fields"]["shipper"]["value"], "Atlas Ltd\n11 Harbor Road")
        self.assertEqual(result["fields"]["consignee"]["value"], "Beacon Ltd\n15 Wharf Street")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_xlsx_third_address_cell_and_formula_warning(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Shipping"
        sheet.append(["BILL OF LADING", "123456"])
        sheet.append(["Shipper", "Atlas Ltd", "11 Harbor Road"])
        sheet.append(["Consignee", "Beacon Ltd"])
        sheet.append(["Gross Weight (KG)", "=100+20"])
        stream = io.BytesIO()
        workbook.save(stream)
        result = read_document(stream.getvalue(), "document.xlsx")
        self.assertIsNone(result["error"])
        self.assertEqual(result["kind"], "BL")
        self.assertEqual(result["fields"]["shipper"]["value"], "Atlas Ltd\n11 Harbor Road")
        self.assertIsNone(result["fields"]["gross_weight_kg"]["value"])
        self.assertTrue(any("Uncalculated formula" in warning for warning in result["warnings"]))

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow unavailable")
    def test_missing_image_ocr_is_an_explicit_failure(self):
        from PIL import Image
        stream = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(stream, format="PNG")
        with patch("harborlight.readers._tesseract", return_value=None):
            result = read_document(stream.getvalue(), "scan.png")
        self.assertIn("Tesseract", result["error"])
        self.assertTrue(read_document(stream.getvalue(), "scan.png", ocr=False)["error"])

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow unavailable")
    def test_ocr_evidence_method_and_confidence(self):
        from PIL import Image
        stream = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(stream, format="PNG")
        with patch("harborlight.readers._ocr_image", return_value=(SI, 0.73, [])):
            result = read_document(stream.getvalue(), "scan.png")
        self.assertIsNone(result["error"])
        self.assertEqual(result["fields"]["shipper"]["confidence"], .73)
        self.assertTrue(result["fields"]["shipper"]["method"].startswith("ocr:"))

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow unavailable")
    def test_low_ocr_confidence_preserves_evidence_but_abstains(self):
        from PIL import Image
        stream = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(stream, format="PNG")
        with patch("harborlight.readers._ocr_image", return_value=(SI, 0.42, [])):
            result = read_document(stream.getvalue(), "scan.png")
        field = result["fields"]["shipper"]
        self.assertIsNone(field["value"])
        self.assertIn("Atlas", field["extracted_value"])
        self.assertIn("Atlas", field["evidence"])
        self.assertTrue(field["uncertain"])
        self.assertTrue(any("Low OCR confidence" in warning for warning in result["warnings"]))

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow unavailable")
    def test_bad_numeric_line_cannot_hide_inside_good_page_average(self):
        from PIL import Image
        text = "BILL OF LADING\nShipper: Atlas Ltd\nGross Weight: 48,200 KG"
        stream = io.BytesIO()
        Image.new("RGB", (100, 100), "white").save(stream, format="PNG")
        with patch("harborlight.readers._ocr_image", return_value=(text, 0.90, [.90, .90, .25])):
            result = read_document(stream.getvalue(), "scan.png")
        self.assertEqual(result["fields"]["shipper"]["value"], "Atlas Ltd")
        self.assertIsNone(result["fields"]["gross_weight_kg"]["value"])
        self.assertEqual(result["fields"]["gross_weight_kg"]["confidence"], .25)

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow unavailable")
    def test_tesseract_tsv_retains_weakest_word_confidence(self):
        from PIL import Image
        from types import SimpleNamespace
        tsv = "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n5\t1\t1\t1\t1\t1\t10\t10\t30\t10\t98\tGross\n5\t1\t1\t1\t1\t2\t45\t10\t40\t10\t96\tWeight:\n5\t1\t1\t1\t1\t3\t90\t10\t40\t10\t35\t48200\n"
        with patch("harborlight.readers._tesseract", return_value="local-tesseract"), patch("harborlight.readers.subprocess.run", return_value=SimpleNamespace(returncode=0, stdout=tsv, stderr="")) as run:
            text, page_confidence, line_confidences = _ocr_image(Image.new("RGB", (150, 100), "white"))
        self.assertEqual(text, "Gross Weight: 48200")
        self.assertGreater(page_confidence, .65)
        self.assertEqual(line_confidences, [.35])
        self.assertNotIn("shell", run.call_args.kwargs)
        self.assertEqual(run.call_args.kwargs["timeout"], 60)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_sparse_workbook_with_excessive_dimensions_is_rejected(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.cell(20000, 40, "sparse")
        stream = io.BytesIO()
        workbook.save(stream)
        result = read_document(stream.getvalue(), "large.xlsx")
        self.assertIn("cell limits", result["error"])

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "pypdf unavailable")
    def test_blank_pdf_routes_to_ocr_requirement(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        stream = io.BytesIO()
        writer.write(stream)
        with patch("harborlight.readers._tesseract", return_value=None):
            result = read_document(stream.getvalue(), "scan.pdf")
        self.assertIn("OCR", result["error"])
        self.assertEqual(result["pages"][0]["method"], "failed")

    @unittest.skipUnless(importlib.util.find_spec("reportlab"), "optional PDF test fixture dependency unavailable")
    def test_digital_pdf_keeps_page_evidence(self):
        from reportlab.pdfgen.canvas import Canvas
        stream = io.BytesIO()
        canvas = Canvas(stream)
        canvas.drawString(50, 780, "SHIPPING INSTRUCTION")
        canvas.drawString(50, 755, "Shipper: Atlas Paper Ltd")
        canvas.drawString(50, 730, "Consignee: Beacon Trading, 15 Wharf Street")
        canvas.showPage()
        canvas.drawString(50, 780, "Total Gross Weight: 48,200 KG")
        canvas.drawString(50, 755, "Container Count: 2 x 40HC")
        canvas.save()
        result = read_document(stream.getvalue(), "file.pdf", ocr=False)
        self.assertIsNone(result["error"])
        self.assertEqual(result["fields"]["gross_weight_kg"]["page"], 2)
        self.assertEqual(result["fields"]["gross_weight_kg"]["value"], "48,200 KG")


if __name__ == "__main__":
    unittest.main()
