"""Rebuild original fictional demo fixtures; needs requirements-dev.txt.

These scenarios are authored independently of organizer datasets and answer keys.
The intentionally broken PDF is a reliability test, not a document to display.
"""
from pathlib import Path
import json
import copy
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent / "demo"
FIELDS = {"shipper": "Shipper", "consignee": "Consignee", "notify_party": "Notify Party", "port_of_loading": "Port of Loading", "port_of_discharge": "Port of Discharge", "container_count": "Total Containers", "gross_weight_kg": "Gross Weight (KG)"}
BASE = {"shipper": "Meridian Paper Industries Ltd", "consignee": "Northstar Trading Pte Ltd", "notify_party": "Northstar Trading Pte Ltd", "port_of_loading": "Port Klang", "port_of_discharge": "Singapore", "container_count": "3", "gross_weight_kg": "22000 kg"}


def text(role, values):
    title = "SHIPPING INSTRUCTION" if role == "SI" else "DRAFT BILL OF LADING"
    return title + "\n" + "\n".join(FIELDS[k] + ": " + str(v) for k, v in values.items()) + "\n"


def document(name, role, values, fmt="txt"):
    filename = f"{name}_{role}.{fmt}"
    path = ROOT / "attachments" / filename
    content = text(role, values)
    if fmt == "txt":
        path.write_text(content, encoding="utf-8")
    elif fmt == "pdf":
        c = canvas.Canvas(str(path), pagesize=(595, 842))
        c.setFillColor(HexColor("#163c43"))
        c.rect(0, 744, 595, 98, fill=1, stroke=0)
        c.setFillColor(HexColor("#ffffff"))
        c.setFont("Helvetica-Bold", 18)
        c.drawString(38, 791, content.splitlines()[0])
        c.setFont("Helvetica", 10)
        c.drawString(38, 767, f"HARBORLIGHT DEMO / {name} / FICTIONAL SHIPMENT")
        y = 690
        for line in content.splitlines()[1:]:
            label, value = line.split(": ", 1)
            c.setFillColor(HexColor("#607577")); c.setFont("Helvetica", 10)
            c.drawString(38, y, label)
            c.setFillColor(HexColor("#173b40")); c.setFont("Helvetica-Bold", 13)
            c.drawString(38, y - 21, value)
            c.setStrokeColor(HexColor("#dfe6e3")); c.line(38, y - 35, 557, y - 35)
            y -= 78
        c.setFont("Helvetica", 9); c.drawString(38, 55, "Synthetic document for a shipping verification demonstration. Page 1 of 1.")
        c.save()
    elif fmt == "docx":
        doc = Document(); doc.add_heading(content.splitlines()[0], 0)
        doc.add_paragraph("Harborlight fictional demonstration shipment")
        table = doc.add_table(rows=0, cols=2)
        for key, value in values.items():
            cells = table.add_row().cells
            cells[0].text, cells[1].text = FIELDS[key], str(value)
        doc.save(path)
    elif fmt == "xlsx":
        book = Workbook(); sheet = book.active; sheet.title = "Shipment"
        sheet.append([content.splitlines()[0]])
        for key, value in values.items():
            sheet.append([FIELDS[key], str(value)])
        sheet.column_dimensions["A"].width = 26; sheet.column_dimensions["B"].width = 48
        book.save(path)
    elif fmt == "png":
        image = Image.new("RGB", (1700, 1200), "white")
        draw = ImageDraw.Draw(image)
        font_path = next((p for p in [Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")] if p.exists()), None)
        font = ImageFont.truetype(str(font_path), 34) if font_path else ImageFont.load_default(size=30)
        for i, line in enumerate(content.splitlines()):
            draw.text((75, 70 + i * 110), line, fill="black", font=font)
        image.save(path)
    return "attachments/" + filename


def email(key, subject, body, paths=None):
    record = {"email_id": key, "from": "operations@example.test", "subject": subject, "body": body, "attachments": paths or []}
    (ROOT / "inbox" / (key + ".json")).write_text(json.dumps(record, indent=2), encoding="utf-8")


def main():
    (ROOT / "inbox").mkdir(parents=True, exist_ok=True)
    (ROOT / "attachments").mkdir(parents=True, exist_ok=True)
    wrong = BASE | {"port_of_discharge": "Bangkok", "container_count": "4"}
    email("HL-001", "Draft BL - closing today", "Please compare the attached draft BL against our shipping instruction. The vessel cut-off is today.", [document("HL-001", "SI", BASE, "pdf"), document("HL-001", "BL", wrong, "pdf")])
    equivalent = BASE | {"shipper": "MERIDIAN PAPER INDUSTRIES LIMITED", "gross_weight_kg": "22 tonnes", "container_count": "3 x 40HC"}
    email("HL-002", "Please verify these shipment documents", "Please check the SI and draft BL. Formatting differs but shipment details should agree.", [document("HL-002", "SI", BASE), document("HL-002", "BL", equivalent)])
    missing = BASE | {"gross_weight_kg": "TBA"}
    email("HL-003", "Check draft before release", "Compare our SI with the draft bill of lading. We need the weight confirmed from operations.", [document("HL-003", "SI", missing, "docx"), document("HL-003", "BL", BASE, "docx")])
    other = ROOT / "attachments" / "HL-004_BL.txt"
    other.write_text("COMMERCIAL INVOICE\nSeller: Meridian Paper Industries\nBuyer: Northstar Trading\nInvoice Total: USD 42,500\n", encoding="utf-8")
    email("HL-004", "SI and BL for comparison", "Verify the attached draft BL against the SI. Please check whether the correct file was attached.", [document("HL-004", "SI", BASE), "attachments/HL-004_BL.txt"])
    email("HL-005", "Draft BL verification", "Please compare the draft BL against the SI. I have attached the instruction; the carrier's draft will follow.", [document("HL-005", "SI", BASE)])
    scan = document("HL-006", "BL", BASE, "png")
    image = Image.open(ROOT / scan)
    scan_pdf = ROOT / "attachments" / "HL-006_BL.pdf"
    image.save(scan_pdf, "PDF", resolution=150)
    email("HL-006", "Scanned BL - verify document", "Please compare this scanned draft BL with the shipping instruction.", [document("HL-006", "SI", BASE), "attachments/HL-006_BL.pdf"])
    email("HL-007", "Prepare shipping instruction", "Please create a new shipping instruction for next week's shipment. Use the booking information we will provide.")
    email("HL-008", "Invoice amount clarification", "Could you explain the detention charges on this invoice? The billed amount differs from our quote.")
    email("HL-009", "Weekly operations notice", "The warehouse will close at 5pm on Friday. This is an operational update; no documents require checking.")
    email("HL-010", "Claim your exclusive prize", "Congratulations! You won a cash prize. Click here to claim and enter your bank details.")
    (ROOT / "attachments" / "HL-011_BL.pdf").write_bytes(b"%PDF-1.7\ncorrupt fixture")
    email("HL-011", "Check the attached draft BL", "Please compare the attached bill of lading with our shipping instruction.", [document("HL-011", "SI", BASE), "attachments/HL-011_BL.pdf"])
    email("HL-012", "Verify draft - weight discrepancy suspected", "Please check the SI and draft BL before finalizing. The gross weight may have been transcribed incorrectly.", [document("HL-012", "SI", BASE, "xlsx"), document("HL-012", "BL", BASE | {"gross_weight_kg": "22500 kg"}, "docx")])
    print("Built 12 fictional cases: digital PDF, DOCX, XLSX, scanned PDF, clean formatting, discrepancies and review failures.")


if __name__ == "__main__":
    main()
