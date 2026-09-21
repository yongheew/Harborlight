"""Local, evidence-preserving shipping document readers.

No filename, answer key, network service, or document instruction determines a
field value. Optional OCR runs the locally installed Tesseract executable.
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from typing import Any

FIELDS = (
    "shipper", "consignee", "notify_party", "port_of_loading",
    "port_of_discharge", "container_count", "gross_weight_kg",
)
MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 100
MAX_TEXT = 2_000_000
MAX_WORKBOOK_CELLS = 250_000
MIN_OCR_CONFIDENCE = 0.65
PARTIES = {"shipper", "consignee", "notify_party"}

# Ordering matters: a long label must consume its short prefix as one label.
_ALIASES = {
    "shipper": r"(?:shipper\s*[/&]\s*exporter|shipper|exporter)(?:\s+(?:name\s*(?:and|&)\s*address|name|address))?",
    "consignee": r"(?:consignee(?:\s+(?:name\s*(?:and|&)\s*address|name|address))?|to\s+(?:the\s+)?order\s+of)",
    "notify_party": r"(?:notify\s+party(?:\s*/\s*intermediate\s+consignee)?|notify)(?:\s+(?:name\s*(?:and|&)\s*address|name|address))?",
    "port_of_loading": r"(?:port\s+of\s+loading|loading\s+port|load\s+port|P\.?O\.?L\.?)",
    "port_of_discharge": r"(?:port\s+of\s+discharge|discharge\s+port|port\s+of\s+unloading|P\.?O\.?D\.?)",
    "container_count": r"(?:(?:total\s+)?(?:no\.?|number)\s*(?:of\s+)?containers?(?:\s+or\s+packages)?|total\s+(?:no\.?\s+of\s+)?containers?|container\s+(?:count|quantity)|containers?\s+(?:count|quantity))",
    "gross_weight_kg": r"(?:(?:total\s+)?gross\s*(?:weight|wt\.?)(?:\s*毛重)?(?:\s*(?:in\s+)?(?:kilograms?|kgs?|metric\s+tonnes?|metric\s+tons?|tonnes?|tons?|mts?|lbs?|pounds?))?)",
}
_LABELS = [
    (key, re.compile(r"(?<![\w])(?:" + pattern + r")(?![\w])(?:[ \t]*(?:\([^\n)]{0,100}\)|\[[^\n\]]{0,70}\]))*[ \t]*(?:[:=][ \t]*)?", re.I))
    for key, pattern in _ALIASES.items()
]
_BOUNDARY = re.compile(
    r"^\s*(?:vessel(?:\s+name)?|ocean\s+vessel|export\s+carrier|voy(?:age|\.)?(?:\s+no\.?)?|"
    r"commodity|description(?:\s+of\s+goods)?|kinds\s+of\s+packages|goods|hs\s*code|"
    r"booking(?:\s+(?:ref|reference|no\.?))?|b\s*/\s*l\b|bill\s+of\s+lading|"
    r"shipping\s+instruction|freight|order\s+no|oc\s*no|container\s+no|seal\s+no|"
    r"net\s+(?:weight|wt)|measurement|place\s+of\s+(?:receipt|delivery)|final\s+destination|"
    r"marks(?:\s+and\s+numbers)?|signature|signed|date|page\s+\d|commercial\s+invoice)\b", re.I,
)
_MISSING = re.compile(r"^(?:[\s_\-–—./]*|n\s*[/.-]?\s*a\.?|none|null|nil|not\s+(?:provided|available|applicable|specified)|missing|blank|tbd|tbc|to\s+be\s+(?:confirmed|advised)|pending|\[(?:blank|missing|redacted)\]|<missing>)$", re.I)


def ocr_available() -> bool:
    """Whether an executable has been configured; does not invoke a process."""
    return _tesseract() is not None


def _tesseract() -> str | None:
    configured = os.environ.get("TESSERACT_CMD", "").strip()
    if configured:
        return shutil.which(configured) or (configured if Path(configured).is_file() else None)
    return shutil.which("tesseract")


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return text.replace("\x00", "").replace("\u200b", "").replace("\ufeff", "").replace("■", " ").replace("\r\n", "\n").replace("\r", "\n")


def _empty_fields() -> dict[str, dict]:
    return {key: {"value": None, "evidence": "", "page": 1, "confidence": 0.0, "method": "not_found"} for key in FIELDS}


def _detect_format(data: bytes, filename: str) -> str:
    if data.lstrip()[:5] == b"%PDF-":
        return "pdf"
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(x.file_size for x in members) > 100 * 1024 * 1024:
                raise ValueError("Office archive exceeds the safe expanded-size limit.")
            names = {x.filename for x in members}
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "xlsx"
        raise ValueError("ZIP attachment is not a supported DOCX or XLSX document.")
    if data.startswith((b"\x89PNG\r\n", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*", b"BM", b"GIF87a", b"GIF89a")) or (data[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        return "image"
    if Path(filename).suffix.lower() in {".txt", ".text", ".csv", ".md", ""}:
        return "text"
    raise ValueError("Unsupported or invalid document content. Supported: TXT, PDF, DOCX, XLSX, PNG, JPEG, TIFF, WEBP.")


def _decode_text(data: bytes, warnings: list[str]) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        warnings.append("Text was decoded as Windows-1252; verify special characters.")
        return text


def _table_lines(rows: list[list[str]]) -> list[str]:
    """Keep adjacent label/value cells together and multiline addresses intact."""
    result: list[str] = []
    column_layout = False
    widths: list[int] = []
    for row in rows:
        for index, cell in enumerate(row):
            while len(widths) <= index:
                widths.append(0)
            widths[index] = max(widths[index], max((len(line) for line in _clean(cell).splitlines()), default=0))
    for row in rows:
        cells = [_clean(cell).strip() for cell in row]
        original_cells = cells[:]
        # Word merged cells are repeated by python-docx; ignore identical neighbors.
        cells = [cell for i, cell in enumerate(cells) if i == 0 or cell != cells[i - 1]]
        if not any(cells):
            result.append("")
            continue
        label_indices = [i for i, cell in enumerate(cells) if _label_only(cell)]
        labels_are_column_headers = len(label_indices) > 1 and len(label_indices) == sum(bool(cell) for cell in cells)
        if labels_are_column_headers:
            column_layout = True
        elif label_indices:
            column_layout = False
        if column_layout:
            # Tables can place party labels over parallel address columns.
            # Keep horizontal positions so their following cells never merge.
            split_cells = [cell.splitlines() for cell in original_cells]
            for line_index in range(max(map(len, split_cells), default=1)):
                result.append("".join((lines[line_index] if line_index < len(lines) else "").ljust(widths[index] + 4) for index, lines in enumerate(split_cells)).rstrip())
        elif label_indices and label_indices[0] == 0:
            for n, index in enumerate(label_indices):
                end = label_indices[n + 1] if n + 1 < len(label_indices) else len(cells)
                value = "\n".join(c for c in cells[index + 1:end] if c)
                result.append(cells[index].rstrip(" :") + ": " + value)
        else:
            result.append("    ".join(cells))
    return result


def _label_only(text: str) -> bool:
    return any(pattern.fullmatch(text.strip()) for _, pattern in _LABELS)


def _read_docx(data: bytes) -> list[dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    # Iterate XML children to retain the actual paragraph/table order.
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            lines.append(Paragraph(element, document).text)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            lines.extend(_table_lines([[cell.text for cell in row.cells] for row in table.rows]))
    # DOCX pagination is not stable until Word renders it. Mark this explicitly.
    return [{"page": 1, "text": "\n".join(lines), "method": "docx:paragraphs+tables", "locator": "document body (logical page)", "confidence": 0.97}]


def _read_xlsx(data: bytes, warnings: list[str]) -> list[dict]:
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    pages: list[dict] = []
    cell_count = 0
    uncached_formulas = 0
    try:
        for index, sheet in enumerate(workbook.worksheets, 1):
            if sheet.max_row is None or sheet.max_column is None:
                raise ValueError("Workbook has no usable worksheet dimensions; open and resave it in Excel before retrying.")
            cell_count += sheet.max_row * sheet.max_column
            if index > MAX_PAGES or sheet.max_row > 20000 or sheet.max_column > 256 or cell_count > MAX_WORKBOOK_CELLS:
                raise ValueError("Workbook exceeds sheet or cell limits.")
            rows = []
            # Walk both streams once. Random access on a read-only worksheet
            # reparses XML and would make many formula cells quadratic in cost.
            for cached_row, formula_row in zip(sheet.iter_rows(), formulas[sheet.title].iter_rows()):
                rows.append(["" if cell.value is None else str(cell.value) for cell in cached_row])
                for cell, cached in zip(formula_row, cached_row):
                    if cell.data_type == "f" and cached.value is None:
                        uncached_formulas += 1
                        if uncached_formulas <= 20:
                            warnings.append(f"Uncalculated formula in sheet {sheet.title}, cell {cell.coordinate}; its value is unavailable.")
            pages.append({"page": index, "text": "\n".join(_table_lines(rows)), "method": "xlsx:cells", "locator": f"sheet: {sheet.title}", "confidence": 0.98})
    finally:
        workbook.close()
        formulas.close()
    if uncached_formulas > 20:
        warnings.append(f"{uncached_formulas - 20} additional uncalculated formula cells have unavailable values.")
    return pages


def _ocr_image(image: Any) -> tuple[str, float, list[float]]:
    executable = _tesseract()
    if not executable:
        raise RuntimeError("OCR required but Tesseract is unavailable. Install Tesseract and set TESSERACT_CMD, then retry.")
    with tempfile.TemporaryDirectory(prefix="harborlight-ocr-") as directory:
        image_path = Path(directory) / "page.png"
        prepared = image.convert("RGB")
        try:
            prepared.thumbnail((4500, 4500))
            prepared.save(image_path)
        finally:
            prepared.close()
        process = subprocess.run(
            [executable, str(image_path), "stdout", "-l", os.environ.get("OCR_LANG", "eng"), "--psm", "6", "tsv"],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60,
            check=False, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if process.returncode:
            raise RuntimeError("Tesseract OCR failed: " + process.stderr.strip()[-400:])
    lines: dict[tuple[str, ...], list[tuple[int, str, int, float]]] = {}
    confidences = []
    for row in csv.DictReader(io.StringIO(process.stdout), delimiter="\t"):
        word = (row.get("text") or "").strip()
        if not word:
            continue
        try:
            confidence = float(row.get("conf", "-1"))
            left = int(row.get("left", "0"))
            width = int(row.get("width", "0"))
        except (ValueError, TypeError):
            continue
        key = tuple(row.get(k, "") for k in ("page_num", "block_num", "par_num", "line_num"))
        lines.setdefault(key, []).append((left, word, width, max(0, confidence / 100)))
        if confidence >= 0:
            confidences.append(confidence / 100)
    rendered = []
    line_confidences = []
    for words in lines.values():
        text = ""
        last_end = 0
        for left, word, width, confidence in sorted(words):
            gap = left - last_end
            text += ("    " if text and gap > 30 else " " if text else "") + word
            last_end = left + width
        rendered.append(text)
        # A low-confidence word in a critical number/address must not disappear
        # inside a high page average. Field extraction takes its weakest line.
        line_confidences.append(min(0.90, min(word[3] for word in words)))
    mean = sum(confidences) / len(confidences) if confidences else 0.0
    return "\n".join(rendered), min(0.90, mean), line_confidences


def _read_image(data: bytes, ocr: bool) -> list[dict]:
    from PIL import Image, ImageSequence
    if not ocr:
        raise RuntimeError("Image document requires OCR, but OCR is disabled.")
    with Image.open(io.BytesIO(data)) as source:
        if source.width * source.height > 40_000_000:
            raise ValueError("Image exceeds the 40-megapixel limit.")
        pages = []
        for index, frame in enumerate(ImageSequence.Iterator(source), 1):
            if index > MAX_PAGES:
                raise ValueError("Image exceeds the page limit.")
            if frame.width * frame.height > 40_000_000:
                raise ValueError("Image frame exceeds the 40-megapixel limit.")
            image = frame.copy()
            try:
                text, confidence, line_confidences = _ocr_image(image)
            finally:
                image.close()
            pages.append({"page": index, "text": text, "method": "ocr:tesseract", "confidence": confidence, "line_confidences": line_confidences})
        return pages


def _read_pdf(data: bytes, ocr: bool, warnings: list[str]) -> tuple[list[dict], str | None]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("Password-protected PDF cannot be read without its password.")
    if len(reader.pages) > MAX_PAGES:
        raise ValueError(f"PDF exceeds the {MAX_PAGES}-page limit.")
    pages = []
    failures = []
    renderer = None
    try:
        for number, page in enumerate(reader.pages, 1):
            try:
                try:
                    text = page.extract_text(extraction_mode="layout") or ""
                except Exception:
                    text = page.extract_text() or ""
                    warnings.append(f"PDF page {number}: layout extraction failed; plain text fallback used.")
                method, confidence = "pdf:text-layout", 0.95
                line_confidences = None
                text_characters = len(re.sub(r"\W", "", text))
                # A short digital continuation page is valid. Sparse text needs
                # OCR only when accompanied by an embedded scan, or truly empty.
                has_image = bool(page.images) if text_characters < 100 else False
                if text_characters < 8 or (text_characters < 100 and has_image):
                    if not ocr:
                        raise RuntimeError("scanned/empty page requires OCR, but OCR is disabled")
                    if not _tesseract():
                        raise RuntimeError("scanned/empty page requires OCR; Tesseract is unavailable (set TESSERACT_CMD)")
                    if renderer is None:
                        import pypdfium2
                        renderer = pypdfium2.PdfDocument(data)
                    rendered_page = renderer[number - 1]
                    try:
                        width, height = rendered_page.get_size()
                        bitmap = rendered_page.render(scale=min(3.0, 4500 / max(width, height)))
                        try:
                            image = bitmap.to_pil()
                            try:
                                text, confidence, line_confidences = _ocr_image(image)
                            finally:
                                image.close()
                        finally:
                            bitmap.close()
                    finally:
                        rendered_page.close()
                    method = "ocr:tesseract"
                pages.append({"page": number, "text": text, "method": method, "confidence": confidence, "line_confidences": line_confidences})
            except Exception as error:
                message = f"PDF page {number}: {error}"
                failures.append(message)
                warnings.append(message)
                pages.append({"page": number, "text": "", "method": "failed", "confidence": 0.0})
    finally:
        if renderer is not None:
            renderer.close()
    return pages, "; ".join(failures) or None


def _kind(text: str) -> tuple[str, str | None]:
    # Content titles only. A filename such as BL.pdf does not establish its type.
    title_lines = [re.sub(r"\s+", " ", x).strip(" :-_=\t") for x in text.splitlines() if x.strip()][:35]
    other = any(re.match(r"^(?:commercial\s+invoice|tax\s+invoice|pro\s*forma\s+invoice|invoice(?:\s+(?:no|number))?\b|packing\s+list|certificate\s+of\s+origin|purchase\s+order)\b", line, re.I) for line in title_lines[:12])
    si = any(re.match(r"^(?:shipping\s+instructions?|(?:bill\s+of\s+lading|b\s*/?\s*l)\s+instructions?)(?:\b|$)", line, re.I) for line in title_lines)
    bl = any(re.match(r"^(?:draft\s+)?(?:bill\s+of\s+lading|b\s*/?\s*l(?:\s+draft)?)(?:\s*\(\s*draft\s*\)|\s+draft)?(?:\s*$|\s*[:#]|\s+no\.?|\s+number|\s+\d)", line, re.I) for line in title_lines)
    if other:
        return "OTHER", "Document content identifies an invoice, packing list, or other unsupported shipping document."
    if si:
        return "SI", None
    if bl:
        return "BL", None
    return "UNKNOWN", "Document type could not be established from its content; filename was not trusted."


def _find_labels(line: str) -> list[tuple[int, int, str, str]]:
    found = []
    for field, pattern in _LABELS:
        for match in pattern.finditer(line):
            prefix = line[:match.start()]
            # Suppress labels embedded in prose or company names. Column breaks,
            # start-of-line and explicit label colons are permitted boundaries.
            if prefix.strip() and not re.search(r"(?:\s{2,}|\t|\|)\s*$", prefix):
                continue
            found.append((match.start(), match.end(), field, match.group()))
    found.sort(key=lambda x: (x[0], -(x[1] - x[0])))
    clean = []
    for item in found:
        if not clean or item[0] >= clean[-1][1]:
            clean.append(item)
    return clean


def _looks_value(field: str, value: str) -> bool:
    if not value.strip():
        return False
    if _MISSING.fullmatch(value.strip()):
        return True
    if field == "gross_weight_kg":
        return bool(re.match(r"^[~≈]?\s*[\d.,]+(?:\s|$|[a-z])", value, re.I))
    if field == "container_count":
        return bool(re.match(r"^(?:\d|one\b|two\b|three\b|four\b|five\b|six\b|seven\b|eight\b|nine\b|ten\b)", value, re.I))
    return True


def _with_label_unit(field: str, value: str, label: str) -> str:
    """Preserve units printed in the label when the value cell is only a number."""
    if field != "gross_weight_kg" or not re.fullmatch(r"[\d\s.,]+", value):
        return value
    unit = re.search(r"\b(metric\s+tonnes?|metric\s+tons?|tonnes?|tons?|mts?|lbs?|pounds?)\b", label, re.I)
    return value + " " + unit.group() if unit else value


def _extract_fields(pages: list[dict], warnings: list[str]) -> dict:
    candidates: dict[str, list[dict]] = {key: [] for key in FIELDS}
    for page in pages:
        active: list[dict] = []

        def flush() -> None:
            nonlocal active
            for item in active:
                value = "\n".join(item["parts"]).strip(" \t|:")
                if not value:
                    continue
                value = _with_label_unit(item["field"], value, item["label"])
                candidates[item["field"]].append({
                    "value": None if _MISSING.fullmatch(value) else value,
                    "evidence": "\n".join(item["evidence"]).strip()[:5000],
                    "page": page["page"], "confidence": item["confidence"],
                    "method": page["method"] + ":label", "locator": page.get("locator", f"page {page['page']}"),
                    "priority": item["priority"],
                })
            active = []

        lines = _clean(page.get("text", "")).expandtabs(4).splitlines()
        for line_index, line in enumerate(lines):
            line_scores = page.get("line_confidences") or []
            line_confidence = line_scores[line_index] if line_index < len(line_scores) else page.get("confidence", 0.95)
            labels = _find_labels(line)
            if labels:
                flush()
                for index, (start, end, field, label) in enumerate(labels):
                    stop = labels[index + 1][0] if index + 1 < len(labels) else len(line)
                    value = line[end:stop].strip(" \t|:")
                    # A numeric table header at the far right has no shipment
                    # total; do not pull a container ID from the next row.
                    table_header = field == "gross_weight_kg" and not value and bool(line[:start].strip())
                    if table_header:
                        continue
                    if value and not _looks_value(field, value):
                        continue
                    active.append({"field": field, "label": label, "parts": [value] if value else [], "evidence": [line[start:stop].strip()], "start": start, "stop": labels[index + 1][0] if index + 1 < len(labels) else None, "confidence": line_confidence, "priority": 2 if re.match(r"\s*total\b", label, re.I) else 1})
                continue
            if not line.strip():
                # Blank lines are not sufficient to terminate an address, but
                # section boundaries and the next label always are.
                continue
            if _BOUNDARY.match(line) or re.match(r"^\s*[-=_]{4,}\s*$", line):
                flush()
                continue
            if not active:
                continue
            for item in active:
                start = item["start"] if len(active) > 1 else 0
                stop = item["stop"]
                continuation = line[start:stop].strip(" \t|")
                if not continuation:
                    continue
                if item["field"] in PARTIES or not item["parts"]:
                    if _looks_value(item["field"], continuation):
                        item["parts"].append(continuation)
                        item["evidence"].append(line[start:stop].rstrip())
                        item["confidence"] = min(item["confidence"], line_confidence)
        flush()

    fields = _empty_fields()
    for field, values in candidates.items():
        if not values:
            continue
        priority = max(value["priority"] for value in values)
        selected = [value for value in values if value["priority"] == priority]
        distinct: dict[str, list[dict]] = {}
        for value in selected:
            canonical = re.sub(r"[^\w]", "", (value["value"] or "").casefold())
            distinct.setdefault(canonical, []).append(value)
        if len(distinct) > 1:
            fields[field] = {"value": None, "evidence": "\n--- conflicting extraction ---\n".join(x["evidence"] for x in selected)[:7000], "page": selected[0]["page"], "confidence": 0.0, "method": "ambiguous", "ambiguous": True, "candidates": [{k: v for k, v in item.items() if k != "priority"} for item in selected]}
            warnings.append(f"Conflicting values for {field}; manual review is required.")
        else:
            winner = max(selected, key=lambda item: item["confidence"])
            fields[field] = {k: v for k, v in winner.items() if k != "priority"}
            if winner["value"] is None:
                fields[field]["confidence"] = 0.0
                fields[field]["method"] += ":placeholder"
            elif winner["method"].startswith("ocr:") and winner["confidence"] < MIN_OCR_CONFIDENCE:
                fields[field]["extracted_value"] = fields[field]["value"]
                fields[field]["value"] = None
                fields[field]["method"] += ":low_confidence"
                fields[field]["uncertain"] = True
                warnings.append(f"Low OCR confidence for {field} ({winner['confidence']:.0%}); its observed text is preserved but requires manual review.")
            if len(selected) > 1:
                fields[field]["occurrences"] = len(selected)
    return fields


def read_document(data: bytes, filename: str, ocr: bool = True) -> dict:
    """Read an attachment; failures are explicit and always JSON-compatible.

    ``page`` is a PDF/image page, XLSX sheet index, or logical DOCX page. Office
    locator metadata disambiguates these without inventing physical pagination.
    Original uploaded bytes are never modified.
    """
    result = {"filename": str(filename), "kind": "UNKNOWN", "text": "", "pages": [], "fields": _empty_fields(), "warnings": [], "error": None, "sha256": hashlib.sha256(data).hexdigest()}
    try:
        if not data:
            raise ValueError("Attachment is empty.")
        if len(data) > MAX_BYTES:
            raise ValueError("Attachment exceeds the 25 MB limit.")
        format_ = _detect_format(data, filename)
        result["format"] = format_
        if format_ == "text":
            text = _decode_text(data, result["warnings"])
            if text.count("\x00") > max(1, len(text) // 100):
                raise ValueError("Text attachment contains binary data.")
            pages = [{"page": 1, "text": text, "method": "text:decode", "confidence": 0.99}]
        elif format_ == "pdf":
            pages, result["error"] = _read_pdf(data, ocr, result["warnings"])
        elif format_ == "docx":
            pages = _read_docx(data)
            result["warnings"].append("DOCX evidence uses logical document order; page 1 is not a rendered Word page number.")
        elif format_ == "xlsx":
            pages = _read_xlsx(data, result["warnings"])
            result["warnings"].append("XLSX evidence page numbers identify worksheet order.")
        else:
            pages = _read_image(data, ocr)
        if sum(len(page["text"]) for page in pages) > MAX_TEXT:
            raise ValueError("Extracted text exceeds the two-million-character limit.")
        for page in pages:
            page["text"] = _clean(page["text"])
        result["pages"] = pages
        result["text"] = "\n\n".join(page["text"] for page in pages)
        if not result["text"].strip():
            result["error"] = result["error"] or "No readable text could be extracted; manual review is required."
        result["kind"], type_warning = _kind(result["text"])
        if type_warning:
            result["warnings"].append(type_warning)
        result["fields"] = _extract_fields(pages, result["warnings"])
        if any(page["method"].startswith("ocr:") for page in pages):
            result["warnings"].append("OCR was used. Confirm low-confidence characters against the source image.")
        missing = [field for field, value in result["fields"].items() if value["value"] is None]
        if missing:
            result["warnings"].append("Missing or uncertain fields: " + ", ".join(missing) + ".")
    except ImportError as error:
        result["error"] = f"Required reader dependency is unavailable: {error}. Install requirements.txt and retry."
    except Exception as error:
        result["error"] = str(error) or type(error).__name__
    return result
